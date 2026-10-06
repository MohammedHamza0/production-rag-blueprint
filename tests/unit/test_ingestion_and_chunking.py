"""Comprehensive unit tests for ingestion normalizers, parsers, and chunkers."""

import io
from pathlib import Path
import pytest
import docx
import pypdf
from src.core.exceptions import DocumentParsingError
from src.domain.protocols import Chunker
from src.ingestion.chunking import (
    ChildChunk,
    ChunkingResult,
    ParentChildChunker,
    ParentChunk,
    SentenceSemanticChunker,
)
from src.ingestion.normalizers import (
    clean_text,
    detect_language_simple,
    estimate_tokens,
    normalize_arabic_letters,
    normalize_for_lexical,
    remove_arabic_diacritics,
    strip_null_bytes,
)
from src.ingestion.parsers import (
    DocxParser,
    MarkdownParser,
    ParsedDocument,
    ParserFactory,
    PDFParser,
    TxtParser,
)


class TestNormalizers:
    """Tests for text cleaning, null-byte removal, and Arabic/Latin normalization."""

    def test_strip_null_bytes(self):
        corrupted = "Text with\x00 null\0 bytes"
        cleaned = strip_null_bytes(corrupted)
        assert "\x00" not in cleaned
        assert "\0" not in cleaned
        assert cleaned == "Text with null bytes"

    def test_clean_text_normalizes_whitespace_and_newlines(self):
        dirty = "  Line 1   with   spaces  \r\n\r\n\r\n\r\nLine 2 \x00 "
        cleaned = clean_text(dirty)
        assert cleaned == "Line 1 with spaces\n\nLine 2"

    def test_remove_arabic_diacritics(self):
        text_with_tashkeel = "بِسْمِ اللَّـهِ الرَّحْمَـٰنِ الرَّحِيمِ تَـمْـدِيـد"
        cleaned = remove_arabic_diacritics(text_with_tashkeel)
        assert "ِ" not in cleaned
        assert "َ" not in cleaned
        assert "ّ" not in cleaned
        assert "ْ" not in cleaned
        assert "ـ" not in cleaned
        assert "بسم" in cleaned

    def test_normalize_arabic_letters(self):
        raw = "إبراهيم وأحمد وآدم وٱبن فاطمة إلى"
        normalized = normalize_arabic_letters(raw, normalize_teh_marbuta=True, normalize_alef_maksura=True)
        assert "إ" not in normalized
        assert "أ" not in normalized
        assert "آ" not in normalized
        assert "ٱ" not in normalized
        assert "فاطمه" in normalized
        assert "الي" in normalized

    def test_normalize_for_lexical(self):
        mixed = "  الْقَاهِرَةُ وَ Alexandria 2026 \x00 "
        lexical = normalize_for_lexical(mixed, lang="ar")
        assert "alexandria" in lexical
        assert "القاهره" in lexical or "القاهرة" in lexical
        assert "\x00" not in lexical

    def test_detect_language_simple(self):
        assert detect_language_simple("هذا نص باللغة العربية الفصحى") == "ar"
        assert detect_language_simple("This is a clean English sentence.") == "en"
        assert detect_language_simple("") == "en"

    def test_estimate_tokens(self):
        assert estimate_tokens("") == 0
        en_toks = estimate_tokens("This is a simple test sentence for token counting.", lang="en")
        ar_toks = estimate_tokens("هذا اختبار بسيط لتقدير عدد الرموز في النص العربي.", lang="ar")
        assert en_toks > 5
        assert ar_toks > 5


class TestTxtParser:
    """Tests for plain text file parsing."""

    @pytest.mark.asyncio
    async def test_parse_txt_file(self, tmp_path: Path):
        file_path = tmp_path / "sample.txt"
        file_path.write_text(
            "First paragraph of the document.\n\nSecond paragraph of the document.",
            encoding="utf-8",
        )

        parser = TxtParser()
        doc = await parser.parse(file_path)

        assert isinstance(doc, ParsedDocument)
        assert doc.extraction_method == "plain-text"
        assert doc.metadata["file_name"] == "sample.txt"
        assert len(doc.sections) >= 1
        assert "First paragraph" in doc.raw_text
        assert "Second paragraph" in doc.raw_text


class TestMarkdownParser:
    """Tests for markdown parsing with structural headings, breadcrumbs, and code protection."""

    @pytest.mark.asyncio
    async def test_parse_markdown_with_hierarchy_and_frontmatter(self, tmp_path: Path):
        md_content = """---
title: Production RAG Guide
author: Staff AI Architect
---

# Architecture Overview
This is the system overview.

## Retrieval Engine
Details about 5-way hybrid search.

```python
# Code block should not be treated as a heading
# Not a heading
def search():
    pass
```

### Reranking Subsystem
Details about cross-encoder reranking.
"""
        file_path = tmp_path / "guide.md"
        file_path.write_text(md_content, encoding="utf-8")

        parser = MarkdownParser()
        doc = await parser.parse(file_path)

        assert doc.metadata.get("title") == "Production RAG Guide"
        assert doc.metadata.get("author") == "Staff AI Architect"
        assert len(doc.sections) == 3

        sec_titles = [s.title for s in doc.sections]
        assert "Architecture Overview" in sec_titles
        assert "Retrieval Engine" in sec_titles
        assert "Reranking Subsystem" in sec_titles

        # Check breadcrumbs on H3 section
        h3_sec = [s for s in doc.sections if s.title == "Reranking Subsystem"][0]
        assert h3_sec.breadcrumbs == ["Architecture Overview", "Retrieval Engine", "Reranking Subsystem"]


class TestDocxParser:
    """Tests for Word document parsing and table-to-markdown conversion."""

    @pytest.mark.asyncio
    async def test_parse_docx_with_headings_and_table(self, tmp_path: Path):
        doc = docx.Document()
        doc.add_heading("System Specifications", level=1)
        doc.add_paragraph("This document outlines database parameters.")

        doc.add_heading("Storage Parameters", level=2)
        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Parameter"
        table.rows[0].cells[1].text = "Value"
        table.rows[1].cells[0].text = "HNSW M"
        table.rows[1].cells[1].text = "16"

        docx_path = tmp_path / "spec.docx"
        doc.save(str(docx_path))

        parser = DocxParser()
        parsed_doc = await parser.parse(docx_path)

        assert parsed_doc.extraction_method == "python-docx"
        assert len(parsed_doc.sections) >= 2
        assert "System Specifications" in parsed_doc.raw_text
        assert "| Parameter | Value |" in parsed_doc.raw_text
        assert "| HNSW M | 16 |" in parsed_doc.raw_text


class TestPDFParser:
    """Tests for PDF text extraction and scanned detection."""

    @pytest.mark.asyncio
    async def test_parse_digital_pdf(self, tmp_path: Path):
        # Generate a minimal valid PDF using pypdf
        writer = pypdf.PdfWriter()
        page = writer.add_blank_page(width=595, height=842)

        # In pypdf, creating text on a blank page can be done by creating an annotation or stream
        # Or using a simple synthetic PDF
        pdf_path = tmp_path / "sample.pdf"
        with open(pdf_path, "wb") as f:
            writer.write(f)

        parser = PDFParser(min_chars_per_page=10)
        parsed = await parser.parse(pdf_path)

        assert isinstance(parsed, ParsedDocument)
        assert parsed.extraction_method == "pypdf"
        assert len(parsed.pages) == 1
        # Blank page should trigger scanned/image-only warning
        assert parsed.is_scanned is True
        assert parsed.has_warnings is True


class TestParserFactory:
    """Tests for factory resolution and error handling."""

    def test_factory_extension_routing(self):
        factory = ParserFactory()
        assert isinstance(factory.get_parser_for_path("doc.pdf"), PDFParser)
        assert isinstance(factory.get_parser_for_path("doc.docx"), DocxParser)
        assert isinstance(factory.get_parser_for_path("doc.md"), MarkdownParser)
        assert isinstance(factory.get_parser_for_path("doc.txt"), TxtParser)

    def test_factory_unsupported_format_raises(self):
        factory = ParserFactory()
        with pytest.raises(DocumentParsingError) as exc_info:
            factory.get_parser_for_path("image.png")
        assert "Unsupported file format '.png'" in str(exc_info.value)


class TestSentenceSemanticChunker:
    """Tests for sentence splitting protecting abbreviations in Arabic and English."""

    def test_splits_sentences_safeguarding_abbreviations(self):
        splitter = SentenceSemanticChunker(lang="ar")
        text = (
            "قام د. أحمد بمراجعة البحث. "
            "أوضح أ. محمد النتائج في ص. 45. "
            "بلغت النسبة 3.14 تقريباً! هل تم التحقق؟ نعم تم ذلك."
        )
        sentences = splitter.split_into_sentences(text)

        assert len(sentences) == 5
        assert "د. أحمد" in sentences[0]
        assert "ص. 45" in sentences[1]
        assert "3.14" in sentences[2]
        assert sentences[3] == "هل تم التحقق؟"
        assert sentences[4] == "نعم تم ذلك."

    def test_english_abbreviations_safeguarded(self):
        splitter = SentenceSemanticChunker(lang="en")
        text = "Dr. Smith met with Prof. Jones at 10.30 a.m. E.g. results were verified! Next step follows."
        sentences = splitter.split_into_sentences(text)

        assert len(sentences) == 3
        assert "Dr. Smith" in sentences[0]
        assert "Prof. Jones" in sentences[0]
        assert "E.g." in sentences[1]
        assert sentences[2] == "Next step follows."


class TestParentChildChunker:
    """Tests for hierarchical chunking, header injection, and protocol compliance."""

    def test_chunker_satisfies_protocol(self):
        chunker = ParentChildChunker()
        assert isinstance(chunker, Chunker)

    def test_chunk_parsed_document_with_header_injection(self):
        chunker = ParentChildChunker(
            parent_target_tokens=300,
            child_target_tokens=80,
            child_overlap_tokens=20,
        )

        # Build synthetic parsed document
        from src.ingestion.parsers.base import ParsedSection

        sec1 = ParsedSection(
            title="Database Architecture",
            level=1,
            text=(
                "PostgreSQL pgvector provides HNSW indexing for dense cosine similarity. "
                "The index uses m=16 and ef_construction=64. "
                "Combined with GIN full-text indexes, it achieves 5-way hybrid RRF ranking."
            ),
            breadcrumbs=["Database Architecture"],
        )
        sec2 = ParsedSection(
            title="Semantic Caching",
            level=1,
            text=(
                "The L1 tier utilizes Redis with pipeline key lookups. "
                "The L2 tier uses pgvector with cosine threshold 0.88. "
                "The CacheEquivalenceVerifier protects against false positives."
            ),
            breadcrumbs=["Semantic Caching"],
        )

        doc = ParsedDocument(
            raw_text=sec1.text + "\n\n" + sec2.text,
            sections=[sec1, sec2],
            metadata={"title": "RAG Blueprint Spec"},
        )

        result: ChunkingResult = chunker.chunk_parsed_document(doc)

        assert result.total_parents == 2
        assert result.total_children >= 2

        # Verify child chunk attributes and header injection
        first_child = result.children[0]
        assert isinstance(first_child, ChildChunk)
        assert first_child.parent_index == 0
        assert "[Document: RAG Blueprint Spec | Section: Database Architecture]" in first_child.header
        assert first_child.header in first_child.content_with_header
        assert first_child.content in first_child.content_with_header
        assert first_child.token_count > 0

        # Verify second section child links to parent 1
        second_parent_child = [c for c in result.children if c.parent_index == 1][0]
        assert "Section: Semantic Caching" in second_parent_child.header

    def test_protocol_chunk_document_method(self):
        chunker = ParentChildChunker()
        sample_text = (
            "Sentence one for test chunking. Sentence two providing context. "
            "Sentence three with additional technical information."
        )
        parent_dicts, child_dicts = chunker.chunk_document(
            raw_text=sample_text,
            doc_title="API Reference",
            parent_target_tokens=500,
            child_target_tokens=100,
        )

        assert len(parent_dicts) >= 1
        assert len(child_dicts) >= 1
        assert "parent_index" in parent_dicts[0]
        assert "content_with_header" in child_dicts[0]
        assert "[Document: API Reference" in child_dicts[0]["header"]
